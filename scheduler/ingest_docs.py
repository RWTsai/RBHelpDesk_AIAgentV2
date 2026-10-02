# ingest_docs.py
"""
文件 GraphRAG 匯入程式（To-Be 新增）
---------------------------------------------------
獨立於 Flask 執行（與 Flask 同一台主機）：
    python scheduler/ingest_docs.py --once        跑一輪就結束
    python scheduler/ingest_docs.py --loop        常駐，每 INGEST_INTERVAL_SEC 秒一輪
    python scheduler/ingest_docs.py --doc 12      只重新處理 DocumentID=12

每一輪：
1. 收件（任何一個來源失敗都不影響其他來源）
   - 後台上傳：IT_KB_SourceDoc.Status=Pending（原檔在 FileContent）
   - 資料夾：掃描 DOC_DROP_FOLDER，新檔／hash 變更 → Pending；檔案消失 → Deleting
   - Google Drive：service account 列出 GDRIVE_FOLDER_IDS（含子資料夾）
   - M365 OneDrive／SharePoint：Graph API + delta 查詢（ONEDRIVE_SHARE_URLS）
2. 刪除：Status=Deleting → 清掉 chunk／relation／EntityChunk 後刪除該筆
3. 匯入：Status=Pending → 解析 → 分塊 → 抽圖譜 → 向量化 → 寫 DB → Done（失敗 Failed + ErrorMsg）

風格：單檔完成、不做 class、中文註解；所有參數讀 .env（config.py）。
"""

import os
import sys
import io
import re
import time
import json
import base64
import hashlib
import argparse
from typing import List
from concurrent.futures import ThreadPoolExecutor

# 讓 `python scheduler/ingest_docs.py` 也能 import 專案根目錄的模組
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import requests
import tiktoken
from openai import OpenAI
from pydantic import BaseModel, Field

from config import (
    get_mssql_conn,
    OPENAI_API_KEY,
    OPENAI_MODEL_EXTRACT,
    OPENAI_MODEL_VISION,
    DOC_VISION_ENABLED,
    DOC_VISION_MIN_CHARS,
    DOC_VISION_MAX_IMAGES,
    DOC_VISION_MIN_PIXELS,
    DOC_VISION_PDF_DPI,
    DOC_VISION_PROMPT,
    DOC_DROP_FOLDER,
    DOC_ALLOWED_EXT,
    INGEST_INTERVAL_SEC,
    INGEST_WORKERS,
    CHUNK_TOKENS,
    CHUNK_OVERLAP,
    EXCEL_ROWS_PER_CHUNK,
    ENTITY_DESC_MAX_CHARS,
    GDRIVE_SA_JSON,
    GDRIVE_FOLDER_IDS,
    M365_TENANT_ID,
    M365_CLIENT_ID,
    M365_CLIENT_SECRET,
    ONEDRIVE_SHARE_URLS,
    ONEDRIVE_RECURSIVE,
    DOC_GRAPH_EXTRACTION_PROMPT,
    ENTITY_SUMMARY_PROMPT,
)
from database.doc_graph_client import norm_name, encode_vec, embed_texts
from utils.logger import log_info, log_error, mask_secrets
from utils.openai_compat import parse_structured

client = OpenAI(api_key=OPENAI_API_KEY)
ENC = tiktoken.get_encoding("cl100k_base")  # text-embedding-3 系列使用的編碼


def _ext(name: str) -> str:
    return os.path.splitext(name or "")[1].lower().lstrip(".")


def _allowed(name: str) -> bool:
    return _ext(name) in DOC_ALLOWED_EXT


def _tokens(text: str) -> int:
    return len(ENC.encode(text or "", disallowed_special=()))


# --- 資料庫連線診斷（連不上時給一句人看得懂的話，而不是一整串 traceback）---
# 用 SQL Server 錯誤碼／SQLSTATE 判斷，不用訊息文字：ODBC 的訊息會跟著作業系統語言變。
# 訊息裡的 SQL Server 錯誤碼 → (原因, 要檢查什麼)
# 依序比對，較精確的放前面：資料庫開不起來時訊息會同時出現 4060 和 18456
_DB_CODE_HINTS = (
    ("4060", "資料庫名稱錯誤，或這個帳號沒有該資料庫的權限", "檢查 MSSQL_CONN 的 DATABASE，以及該帳號在這個資料庫的權限"),
    ("40615", "防火牆未放行這台主機的 IP", "在資料庫端的防火牆規則加入這台主機的 IP"),
    ("18452", "帳號不被信任（可能誤用 Windows 驗證）", "檢查 MSSQL_CONN 是否該用 UID／PWD 或 Trusted_Connection=yes"),
    ("18456", "帳號或密碼不正確", "檢查 .env 的 MSSQL_CONN 裡的 UID／PWD"),
)
# SQLSTATE → (原因, 要檢查什麼)
_DB_STATE_HINTS = {
    "28000": ("登入失敗", "檢查 .env 的 MSSQL_CONN 裡的 UID／PWD"),
    "08001": ("連不到資料庫主機", "檢查 MSSQL_CONN 的 SERVER（主機名稱、實例名、連接埠）與網路／防火牆，並確認 SQL Server 已開啟 TCP/IP"),
    "08S01": ("與資料庫的連線中斷", "確認資料庫主機與網路狀態"),
    "08004": ("資料庫拒絕連線", "檢查 MSSQL_CONN 的 DATABASE 與該帳號的權限"),
    "HYT00": ("連線逾時", "確認資料庫主機是否開機、網路是否可達"),
    "HYT01": ("連線逾時", "確認資料庫主機是否開機、網路是否可達"),
    "IM002": ("ODBC 驅動程式名稱不對或未安裝", "檢查 MSSQL_CONN 的 DRIVER，並確認主機已安裝 ODBC Driver for SQL Server"),
    "IM003": ("ODBC 驅動程式無法載入", "重新安裝 Microsoft ODBC Driver 17／18 for SQL Server"),
    "42000": ("資料庫名稱錯誤，或這個帳號沒有該資料庫的權限", "檢查 MSSQL_CONN 的 DATABASE，以及該帳號在這個資料庫的權限"),
}


def _db_error_lines(e) -> list:
    """把 pyodbc 的錯誤轉成「原因／請檢查／詳細」三行；認不出來就只印原始訊息。"""
    raw = mask_secrets(str(e))
    hint = None

    # 1) 先看 SQL Server 錯誤碼（比 SQLSTATE 精確）
    for code, cause, action in _DB_CODE_HINTS:
        if f"({code})" in raw or f"[{code}]" in raw:
            hint = (cause, action)
            break

    # 2) 再看 SQLSTATE（pyodbc 放在 args[0]）
    if hint is None:
        args = getattr(e, "args", ())
        state = args[0].upper() if args and isinstance(args[0], str) else ""
        hint = _DB_STATE_HINTS.get(state)

    detail = raw if len(raw) <= 300 else raw[:300] + "…"
    if hint is None:
        return [f"原因：{detail}", "請檢查：.env 的 MSSQL_CONN 設定與資料庫主機狀態"]
    return [f"原因：{hint[0]}", f"請檢查：{hint[1]}", f"詳細：{detail}"]


def check_db() -> bool:
    """開工前先確認資料庫連得上。連不上就印出原因並回 False（不丟 traceback）。"""
    try:
        with get_mssql_conn() as conn:
            conn.cursor().execute("SELECT 1").fetchone()
        return True
    except Exception as e:
        log_error("[ingest] 資料庫連線失敗，這一輪不執行")
        for line in _db_error_lines(e):
            log_error(f"[ingest] {line}")
        return False


# =====================================================
# 1. 解析：bytes → [(章節路徑, 文字)]
# =====================================================
def parse_file(file_name: str, data: bytes):
    ext = _ext(file_name)
    if ext == "pdf":
        sections = _parse_pdf(data)
    elif ext == "docx":
        sections = _parse_docx(data)
    elif ext == "xlsx":
        sections = _parse_xlsx(data)
    elif ext == "pptx":
        sections = _parse_pptx(data)
    elif ext == "md":
        sections = _parse_md(_decode_text(data))
    elif ext == "txt":
        sections = [("", _decode_text(data))]
    else:
        raise ValueError(f"不支援的副檔名：{ext}")

    if ext in ("pdf", "pptx"):
        sections = _vision_enrich(ext, data, sections)
    return [(sec, text) for sec, text in sections if text.strip()]


def _decode_text(data: bytes) -> str:
    """MD／TXT 編碼：先試 UTF-8，再試 Big5（舊文件常見）"""
    for enc in ("utf-8-sig", "cp950", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data):
    import fitz  # PyMuPDF
    out = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, 1):
            # 沒有文字的頁面也保留（掃描版），等視覺解析補；補不到最後會被濾掉
            out.append((f"第{i}頁", page.get_text("text").strip()))
    return out


def _parse_docx(data):
    """依文件順序讀段落與表格，Heading／標題 樣式建立章節階層"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    out, heads, buf = [], [], []

    def flush():
        text = "\n".join(buf).strip()
        if text:
            out.append((" > ".join(heads), text))
        buf.clear()

    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"^(Heading|標題)\s*(\d+)", style)
            if m and text:
                flush()
                level = int(m.group(2))
                heads[:] = heads[: level - 1] + [text]
            elif text:
                buf.append(text)
        elif tag == "tbl":
            t = Table(el, d)
            rows = []
            for r in t.rows:
                cells = [c.text.strip().replace("\n", " ") for c in r.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                buf.append("\n".join(rows))
    flush()
    return out


def _parse_xlsx(data):
    """每個工作表轉 Markdown 表格，每 EXCEL_ROWS_PER_CHUNK 列一段（每段都帶表頭）"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ").strip() for v in r]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        header, body = rows[0], rows[1:]
        width = len(header)
        head_md = "| " + " | ".join(header) + " |\n|" + "---|" * width
        if not body:
            out.append((f"工作表：{ws.title}", head_md))
            continue
        # 每段同時受「列數上限」與「token 上限」限制，確保每段都帶表頭、不會在表格中間被切斷
        budget = CHUNK_TOKENS - _tokens(head_md) - 20
        part, part_tok, start = [], 0, 0
        for idx, r in enumerate(body):
            line = "| " + " | ".join((r + [""] * width)[:width]) + " |"
            n = _tokens(line)
            if part and (len(part) >= EXCEL_ROWS_PER_CHUNK or part_tok + n > budget):
                out.append((f"工作表：{ws.title} 第{start + 2}-{start + 1 + len(part)}列", head_md + "\n" + "\n".join(part)))
                part, part_tok, start = [], 0, idx
            part.append(line)
            part_tok += n
        if part:
            out.append((f"工作表：{ws.title} 第{start + 2}-{start + 1 + len(part)}列", head_md + "\n" + "\n".join(part)))
    wb.close()
    return out


def _parse_pptx(data):
    """逐張投影片：標題、文字方塊、表格、備註"""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts, title = [], ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t and t != title:
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows:
                    texts.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[備註] {notes}")
        body = "\n".join(texts).strip()
        if title or body:
            section = f"投影片{i}" + (f"：{title}" if title else "")
            out.append((section, (title + "\n" + body).strip()))
    return out


# =====================================================
# 1b. 視覺解析：文字太少的頁面／投影片，把圖片轉成文字
# =====================================================
_VISION_CACHE = {}          # 圖片 sha256 → 轉出來的文字（同一輪內重複的圖不重付）


def _pptx_images(data):
    """{投影片編號: [圖片 bytes]}；群組裡的圖片也會取出"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = {}

    def walk(shapes, bucket):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, bucket)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                except Exception:
                    continue
                # 太小的圖（logo、項目符號圖示）跳過
                if (sh.width or 0) * (sh.height or 0) and len(blob) < 8000:
                    continue
                bucket.append(blob)

    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        bucket = []
        walk(slide.shapes, bucket)
        if bucket:
            out[i] = bucket
    return out


def _pdf_page_images(data, pages):
    """{頁碼: [PNG bytes]}；整頁算繪成圖（掃描版、向量圖都吃得到）"""
    import fitz
    out = {}
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i in pages:
            if i < 1 or i > doc.page_count:
                continue
            pix = doc[i - 1].get_pixmap(dpi=DOC_VISION_PDF_DPI)
            if pix.width * pix.height < DOC_VISION_MIN_PIXELS:
                continue
            out[i] = [pix.tobytes("png")]
    return out


# 截圖常帶測試帳號與個資；提示詞會要求不要轉錄，這裡再做一次程式端過濾（樣式明確的才擋）
_PII_PATTERNS = [
    (re.compile(r"[\w.\-]+@[\w.\-]+\.\w{2,}"), "（電子郵件）"),
    (re.compile(r"(?<!\d)09\d{2}[-\s]?\d{3}[-\s]?\d{3}(?!\d)"), "（手機）"),
    (re.compile(r"(?<!\d)0[2-8][-\s)]?\d{3,4}[-\s]?\d{4}(?!\d)"), "（電話）"),
    (re.compile(r"(?<![A-Za-z0-9])[A-Z][12]\d{8}(?![A-Za-z0-9])"), "（身分證號）"),
]


def _scrub_pii(text: str) -> str:
    for pattern, replacement in _PII_PATTERNS:
        text = pattern.sub(replacement, text)
    return text


def _vision_text(images, title):
    """把圖片送視覺模型，回傳轉錄的文字"""
    content = [{"type": "text", "text": DOC_VISION_PROMPT.format(title=title or "（無標題）")}]
    for blob in images:
        b64 = base64.b64encode(blob).decode("ascii")
        content.append({"type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{b64}", "detail": "high"}})
    res = client.chat.completions.create(
        model=OPENAI_MODEL_VISION,
        messages=[{"role": "user", "content": content}],
        temperature=0,
    )
    text = (res.choices[0].message.content or "").strip()
    return "" if text in ("（無內容）", "(無內容)") else _scrub_pii(text)


def _vision_enrich(ext, data, sections):
    """
    文字量低於 DOC_VISION_MIN_CHARS 的頁面／投影片才送視覺模型，
    轉出來的文字以 [圖片內容] 標記併回該段，後面的分塊、向量、圖譜流程都吃得到。
    """
    if not DOC_VISION_ENABLED:
        return sections

    # 章節路徑 → 頁碼（第N頁 / 投影片N）
    thin = {}
    for idx, (sec, text) in enumerate(sections):
        if len(text.strip()) >= DOC_VISION_MIN_CHARS:
            continue
        m = re.search(r"(?:第|投影片)(\d+)", sec or "")
        if m:
            thin[int(m.group(1))] = idx
    if not thin:
        return sections

    try:
        images = _pptx_images(data) if ext == "pptx" else _pdf_page_images(data, sorted(thin))
    except Exception as e:
        log_error(f"[ingest] 取圖片失敗，略過視覺解析：{e}")
        return sections

    out = list(sections)
    done = 0
    for page, idx in sorted(thin.items()):
        blobs = images.get(page) or []
        if not blobs:
            continue
        blobs = blobs[:DOC_VISION_MAX_IMAGES]
        key = hashlib.sha256(b"".join(blobs)).hexdigest()
        sec, text = out[idx]
        try:
            if key not in _VISION_CACHE:
                _VISION_CACHE[key] = _vision_text(blobs, sec)
            extra = _VISION_CACHE[key]
        except Exception as e:
            log_error(f"[ingest] 視覺解析失敗（{sec}）：{e}")
            continue
        if extra:
            out[idx] = (sec, (text + "\n[圖片內容]\n" + extra).strip())
            done += 1
    if done:
        log_info(f"[ingest] 視覺解析：{done}/{len(thin)} 頁補上圖片內容")
    return out


def _parse_md(text):
    """依 # 標題切段"""
    out, heads, buf = [], [], []
    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            if "".join(buf).strip():
                out.append((" > ".join(heads), "\n".join(buf).strip()))
            buf = []
            level = len(m.group(1))
            heads[:] = heads[: level - 1] + [m.group(2).strip()]
        else:
            buf.append(line)
    if "".join(buf).strip():
        out.append((" > ".join(heads), "\n".join(buf).strip()))
    return out


# =====================================================
# 2. 分塊：合併過短的段落、切開過長的段落
# =====================================================
def make_chunks(sections):
    """回傳 [(章節路徑, 文字, token 數)]"""
    chunks = []
    step = max(1, CHUNK_TOKENS - CHUNK_OVERLAP)

    # 先把每段切到不超過 CHUNK_TOKENS
    pieces = []
    for section, text in sections:
        text = re.sub(r"[ \t]+", " ", text).strip()
        if len(text) < 2:
            continue
        toks = ENC.encode(text, disallowed_special=())
        if len(toks) <= CHUNK_TOKENS:
            pieces.append((section, text, len(toks)))
        else:
            for start in range(0, len(toks), step):
                part = toks[start:start + CHUNK_TOKENS]
                pieces.append((section, ENC.decode(part), len(part)))
                if start + CHUNK_TOKENS >= len(toks):
                    break

    # 再把相鄰的短段合併（例如很多小標題、很短的頁面）
    cur_sec, cur_texts, cur_tok = [], [], 0
    for section, text, n in pieces:
        if cur_texts and cur_tok + n > CHUNK_TOKENS:
            chunks.append((_merge_section(cur_sec), "\n\n".join(cur_texts), cur_tok))
            cur_sec, cur_texts, cur_tok = [], [], 0
        cur_sec.append(section)
        cur_texts.append(text if len(cur_texts) == 0 or not section else f"【{section}】\n{text}")
        cur_tok += n
    if cur_texts:
        chunks.append((_merge_section(cur_sec), "\n\n".join(cur_texts), cur_tok))
    return chunks


def _merge_section(secs):
    secs = [s for s in secs if s]
    if not secs:
        return ""
    uniq = list(dict.fromkeys(secs))
    return uniq[0] if len(uniq) == 1 else f"{uniq[0]} ～ {uniq[-1]}"


# =====================================================
# 3. 抽取圖譜（Structured Output，沿用 ai_text_node 的 parse 寫法）
# =====================================================
class ExEntity(BaseModel):
    name: str = Field(description="實體名稱")
    type: str = Field(description="實體類型")
    description: str = Field(description="實體說明")


class ExRelation(BaseModel):
    source: str = Field(description="來源實體名稱")
    target: str = Field(description="目標實體名稱")
    relation: str = Field(description="關係")
    description: str = Field(description="關係說明")


class GraphExtraction(BaseModel):
    entities: List[ExEntity] = Field(description="實體清單")
    relations: List[ExRelation] = Field(description="關係清單")


def extract_graph(doc_name, section, text):
    """單一 chunk 抽實體與關係；失敗回傳空結果（不讓整份文件失敗）"""
    try:
        res = parse_structured(
            client,
            model=OPENAI_MODEL_EXTRACT,
            messages=[{"role": "user", "content": DOC_GRAPH_EXTRACTION_PROMPT.format(
                doc_name=doc_name, section=section or "（無）", text=text)}],
            response_format=GraphExtraction,
        )
        return res.choices[0].message.parsed or GraphExtraction(entities=[], relations=[])
    except Exception as e:
        log_error(f"[ingest] 圖譜抽取失敗（{doc_name} / {section}）：{e}")
        return GraphExtraction(entities=[], relations=[])


def _summarize_entity(name, desc):
    try:
        res = client.chat.completions.create(
            model=OPENAI_MODEL_EXTRACT,
            messages=[{"role": "user", "content": ENTITY_SUMMARY_PROMPT.format(name=name, descriptions=desc)}],
        )
        return (res.choices[0].message.content or desc).strip()[:ENTITY_DESC_MAX_CHARS]
    except Exception as e:
        log_error(f"[ingest] 實體描述摘要失敗（{name}）：{e}")
        return desc[:ENTITY_DESC_MAX_CHARS]


def _embed_batched(texts, batch=64):
    out = []
    for i in range(0, len(texts), batch):
        out.extend(embed_texts(texts[i:i + batch]))
    return out


# =====================================================
# 4. 處理單一文件
# =====================================================
def ingest_document(doc_id, file_name, data):
    t0 = time.time()
    sections = parse_file(file_name, data)
    chunks = make_chunks(sections)
    if not chunks:
        raise ValueError("沒有可擷取的文字（可能是掃描檔或空白文件）")
    log_info(f"[ingest] #{doc_id} {file_name}：{len(sections)} 段 → {len(chunks)} 個 chunk")

    # 4.1 並行抽圖譜
    with ThreadPoolExecutor(max_workers=max(1, INGEST_WORKERS)) as pool:
        graphs = list(pool.map(lambda c: extract_graph(file_name, c[0], c[1]), chunks))

    # 4.2 chunk 向量（帶文件名與章節，提高檢索準確度）
    chunk_vecs = _embed_batched([f"{file_name}｜{sec}\n{text}" for sec, text, _ in chunks])

    # 4.3 彙整實體（以 NormName 合併）與關係
    ents = {}   # norm → {name, type, descs[], chunk_idx:set}
    rels = []   # (src_norm, dst_norm, relation, desc, chunk_idx)
    for ci, g in enumerate(graphs):
        for e in g.entities:
            n = norm_name(e.name)
            if not n:
                continue
            item = ents.setdefault(n, {"name": e.name.strip()[:400], "type": e.type[:50], "descs": [], "chunks": set()})
            if e.description and e.description not in item["descs"]:
                item["descs"].append(e.description.strip())
            item["chunks"].add(ci)
        for r in g.relations:
            s, t = norm_name(r.source), norm_name(r.target)
            if not s or not t or s == t or not r.relation.strip():
                continue
            for n, raw in ((s, r.source), (t, r.target)):
                item = ents.setdefault(n, {"name": raw.strip()[:400], "type": "Other", "descs": [], "chunks": set()})
                item["chunks"].add(ci)
            rels.append((s, t, r.relation.strip()[:200], (r.description or "")[:1000], ci))

    # 4.4 先讀既有實體描述，合併後（太長就摘要）重新算向量 —— 在 transaction 外做，避免長時間鎖表
    existing = {}
    if ents:
        with get_mssql_conn() as conn:
            cursor = conn.cursor()
            norms = list(ents.keys())
            for i in range(0, len(norms), 500):
                part = norms[i:i + 500]
                cursor.execute(
                    f"SELECT EntityID, NormName, EntityType, Description FROM IT_KB_Entity "
                    f"WHERE NormName IN ({','.join('?' * len(part))})", part)
                for r in cursor.fetchall():
                    existing[r.NormName] = {"id": r.EntityID, "type": r.EntityType, "desc": r.Description or ""}

    for n, item in ents.items():
        old = existing.get(n, {}).get("desc", "")
        new_parts = [d for d in item["descs"] if d not in old]
        desc = "\n".join([p for p in [old] + new_parts if p]).strip()
        if len(desc) > ENTITY_DESC_MAX_CHARS:
            desc = _summarize_entity(item["name"], desc)
        item["desc"] = desc
    ent_norms = list(ents.keys())
    ent_vecs = _embed_batched([f"{ents[n]['name']}（{ents[n]['type']}）：{ents[n]['desc']}" for n in ent_norms]) if ent_norms else []

    # 4.5 寫入（同一個 transaction）
    conn = get_mssql_conn()
    try:
        conn.autocommit = False
        cursor = conn.cursor()
        _delete_doc_rows(cursor, doc_id)

        chunk_ids = []
        for seq, ((sec, text, ntok), vec) in enumerate(zip(chunks, chunk_vecs)):
            cursor.execute(
                "INSERT INTO IT_KB_Chunk (DocumentID, Seq, SectionPath, ChunkText, Embedding, TokenCount) "
                "OUTPUT INSERTED.ChunkID VALUES (?, ?, ?, ?, ?, ?)",
                (doc_id, seq, sec[:1000], text, encode_vec(vec), ntok))
            chunk_ids.append(cursor.fetchone()[0])

        ent_ids = {}
        for n, vec in zip(ent_norms, ent_vecs):
            item = ents[n]
            if n in existing:
                eid = existing[n]["id"]
                cursor.execute(
                    "UPDATE IT_KB_Entity SET Description = ?, Embedding = ?, "
                    "EntityType = COALESCE(NULLIF(EntityType, ''), ?), UpdatedAt = SYSUTCDATETIME() WHERE EntityID = ?",
                    (item["desc"], encode_vec(vec), item["type"], eid))
            else:
                cursor.execute(
                    "INSERT INTO IT_KB_Entity (Name, NormName, EntityType, Description, Embedding) "
                    "OUTPUT INSERTED.EntityID VALUES (?, ?, ?, ?, ?)",
                    (item["name"], n, item["type"], item["desc"], encode_vec(vec)))
                eid = cursor.fetchone()[0]
            ent_ids[n] = eid
            for ci in item["chunks"]:
                cursor.execute(
                    "INSERT INTO IT_KB_EntityChunk (EntityID, ChunkID, DocumentID) VALUES (?, ?, ?)",
                    (eid, chunk_ids[ci], doc_id))

        seen = set()
        for s, t, rel, desc, ci in rels:
            key = (s, t, rel, ci)
            if key in seen:
                continue
            seen.add(key)
            cursor.execute(
                "INSERT INTO IT_KB_Relation (SrcEntityID, DstEntityID, Relation, Description, ChunkID, DocumentID) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (ent_ids[s], ent_ids[t], rel, desc, chunk_ids[ci], doc_id))

        cursor.execute(
            "UPDATE IT_KB_SourceDoc SET Status = 'Done', ChunkCount = ?, ErrorMsg = NULL, UpdatedAt = SYSUTCDATETIME() "
            "WHERE DocumentID = ?", (len(chunk_ids), doc_id))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    _cleanup_orphan_entities()
    log_info(f"[ingest] #{doc_id} 完成：chunk {len(chunks)}、實體 {len(ents)}、關係 {len(seen)}，耗時 {time.time() - t0:.1f}s")


def _delete_doc_rows(cursor, doc_id):
    cursor.execute("DELETE FROM IT_KB_EntityChunk WHERE DocumentID = ?", doc_id)
    cursor.execute("DELETE FROM IT_KB_Relation WHERE DocumentID = ?", doc_id)
    cursor.execute("DELETE FROM IT_KB_Chunk WHERE DocumentID = ?", doc_id)


def _cleanup_orphan_entities():
    """刪除已經沒有任何 chunk、也沒有任何關係的實體"""
    try:
        with get_mssql_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                DELETE e FROM IT_KB_Entity e
                WHERE NOT EXISTS (SELECT 1 FROM IT_KB_EntityChunk ec WHERE ec.EntityID = e.EntityID)
                  AND NOT EXISTS (SELECT 1 FROM IT_KB_Relation r WHERE r.SrcEntityID = e.EntityID OR r.DstEntityID = e.EntityID)
            """)
            conn.commit()
    except Exception as e:
        log_error(f"[ingest] 清理孤立實體失敗：{e}")


# =====================================================
# 5. IT_KB_SourceDoc 共用操作
# =====================================================
def upsert_source(source_type, source_ref, file_name, file_hash, source_url=None):
    """新檔 → Pending；hash 變更 → Pending；其他只更新名稱／連結"""
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT DocumentID, FileHash, Status FROM IT_KB_SourceDoc WHERE SourceType = ? AND SourceRef = ?",
            (source_type, source_ref))
        row = cursor.fetchone()
        if row is None:
            cursor.execute(
                "INSERT INTO IT_KB_SourceDoc (FileName, SourceType, SourceRef, SourceUrl, FileHash, Status) "
                "VALUES (?, ?, ?, ?, ?, 'Pending')",
                (file_name[:500], source_type, source_ref, source_url, file_hash))
            log_info(f"[ingest] 新文件（{source_type}）：{file_name}")
        elif row.FileHash != file_hash or row.Status == "Deleting":
            cursor.execute(
                "UPDATE IT_KB_SourceDoc SET FileName = ?, SourceUrl = ?, FileHash = ?, Status = 'Pending', "
                "ErrorMsg = NULL, UpdatedAt = SYSUTCDATETIME() WHERE DocumentID = ?",
                (file_name[:500], source_url, file_hash, row.DocumentID))
            log_info(f"[ingest] 文件有變更（{source_type}）：{file_name}")
        else:
            cursor.execute(
                "UPDATE IT_KB_SourceDoc SET FileName = ?, SourceUrl = ? WHERE DocumentID = ? "
                "AND (FileName <> ? OR ISNULL(SourceUrl, '') <> ISNULL(?, ''))",
                (file_name[:500], source_url, row.DocumentID, file_name[:500], source_url))
        conn.commit()


def mark_deleting(source_type, source_ref):
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE IT_KB_SourceDoc SET Status = 'Deleting', UpdatedAt = SYSUTCDATETIME() "
            "WHERE SourceType = ? AND SourceRef = ? AND Status <> 'Deleting'", (source_type, source_ref))
        if cursor.rowcount:
            log_info(f"[ingest] 來源已刪除（{source_type}）：{source_ref}")
        conn.commit()


def _refs_of(source_type):
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT SourceRef FROM IT_KB_SourceDoc WHERE SourceType = ?", source_type)
        return {r.SourceRef for r in cursor.fetchall()}


def _get_sync_state(key):
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT DeltaLink FROM IT_KB_SyncState WHERE SourceKey = ?", key)
        row = cursor.fetchone()
        return row.DeltaLink if row else None


def _save_sync_state(key, delta_link=None, error=None):
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            MERGE IT_KB_SyncState AS T
            USING (SELECT ? AS SourceKey) AS S ON T.SourceKey = S.SourceKey
            WHEN MATCHED THEN UPDATE SET
                DeltaLink = COALESCE(?, T.DeltaLink), LastSyncAt = SYSUTCDATETIME(), LastError = ?
            WHEN NOT MATCHED THEN INSERT (SourceKey, DeltaLink, LastSyncAt, LastError)
                VALUES (?, ?, SYSUTCDATETIME(), ?);
        """, (key, delta_link, error, key, delta_link, error))
        conn.commit()


# =====================================================
# 6. 來源：本機資料夾
# =====================================================
def scan_folder():
    if not DOC_DROP_FOLDER:
        return
    if not os.path.isdir(DOC_DROP_FOLDER):
        log_error(f"[ingest] DOC_DROP_FOLDER 不存在：{DOC_DROP_FOLDER}")
        return
    seen = set()
    for root, _, files in os.walk(DOC_DROP_FOLDER):
        for f in files:
            if f.startswith("~$") or not _allowed(f):  # ~$ 是 Office 暫存檔
                continue
            path = os.path.join(root, f)
            rel = os.path.relpath(path, DOC_DROP_FOLDER).replace("\\", "/")
            with open(path, "rb") as fp:
                h = hashlib.sha256(fp.read()).hexdigest()
            upsert_source("folder", rel, f, h)
            seen.add(rel)

    known = _refs_of("folder")
    # 整個資料夾突然一個檔案都掃不到（網路磁碟掉線、權限被改）時不做刪除，避免整批知識被清掉
    if not seen and known:
        log_error(f"[ingest] {DOC_DROP_FOLDER} 掃不到任何檔案，但先前有 {len(known)} 份文件；"
                  f"這一輪不做刪除，請確認路徑與權限")
        return
    for ref in known - seen:
        mark_deleting("folder", ref)


# =====================================================
# 7. 來源：Google Drive（service account）
# =====================================================
_GOOGLE_EXPORT = {
    "application/vnd.google-apps.document":
        ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "application/vnd.google-apps.spreadsheet":
        ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "application/vnd.google-apps.presentation":
        ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
}


def _gdrive_service():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    creds = service_account.Credentials.from_service_account_file(
        GDRIVE_SA_JSON, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def sync_gdrive():
    if not GDRIVE_FOLDER_IDS:
        return
    if not GDRIVE_SA_JSON or not os.path.isfile(GDRIVE_SA_JSON):
        log_error(f"[ingest] 找不到 GDRIVE_SA_JSON：{GDRIVE_SA_JSON}")
        return

    svc = _gdrive_service()
    seen, all_ok = set(), True
    for folder_id in GDRIVE_FOLDER_IDS:
        key = f"gdrive:{folder_id}"
        try:
            queue = [folder_id]
            while queue:
                fid = queue.pop()
                token = None
                while True:
                    res = svc.files().list(
                        q=f"'{fid}' in parents and trashed = false",
                        fields="nextPageToken, files(id, name, mimeType, modifiedTime, md5Checksum, webViewLink)",
                        pageSize=200, pageToken=token,
                        supportsAllDrives=True, includeItemsFromAllDrives=True,
                    ).execute()
                    for f in res.get("files", []):
                        mime = f["mimeType"]
                        if mime == "application/vnd.google-apps.folder":
                            queue.append(f["id"])
                            continue
                        name = f["name"]
                        if mime in _GOOGLE_EXPORT:
                            name = f"{name}.{_GOOGLE_EXPORT[mime][1]}"
                        if not _allowed(name):
                            continue
                        h = f.get("md5Checksum") or f.get("modifiedTime")
                        upsert_source("gdrive", f["id"], name, h, f.get("webViewLink"))
                        seen.add(f["id"])
                    token = res.get("nextPageToken")
                    if not token:
                        break
            _save_sync_state(key)
        except Exception as e:
            all_ok = False
            log_error(f"[ingest] Google Drive 資料夾 {folder_id} 同步失敗：{e}")
            _save_sync_state(key, error=str(e)[:4000])

    # 所有資料夾都列成功才判斷刪除，避免暫時性錯誤把文件誤刪
    if all_ok:
        for ref in _refs_of("gdrive") - seen:
            mark_deleting("gdrive", ref)


def _fetch_gdrive(file_id):
    from googleapiclient.http import MediaIoBaseDownload
    svc = _gdrive_service()
    meta = svc.files().get(fileId=file_id, fields="mimeType", supportsAllDrives=True).execute()
    if meta["mimeType"] in _GOOGLE_EXPORT:
        req = svc.files().export_media(fileId=file_id, mimeType=_GOOGLE_EXPORT[meta["mimeType"]][0])
    else:
        req = svc.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    return buf.getvalue()


# =====================================================
# 8. 來源：M365 OneDrive / SharePoint（Microsoft Graph）
# =====================================================
GRAPH = "https://graph.microsoft.com/v1.0"
_MSAL = {"app": None}


def _graph_token():
    import msal
    if _MSAL["app"] is None:
        _MSAL["app"] = msal.ConfidentialClientApplication(
            M365_CLIENT_ID,
            authority=f"https://login.microsoftonline.com/{M365_TENANT_ID}",
            client_credential=M365_CLIENT_SECRET,
        )
    r = _MSAL["app"].acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in r:
        raise RuntimeError(f"取得 Graph token 失敗：{r.get('error')} {r.get('error_description')}")
    return r["access_token"]


def _graph_get(url, stream=False):
    full = url if url.startswith("http") else GRAPH + url
    res = requests.get(full, headers={"Authorization": f"Bearer {_graph_token()}"}, timeout=120, stream=stream)
    res.raise_for_status()
    return res


def _share_id(url):
    """分享連結 → Graph /shares/{id}（u! + base64url，去掉補位 =）"""
    return "u!" + base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii").rstrip("=")


def _onedrive_file_hash(item):
    hashes = (item.get("file") or {}).get("hashes") or {}
    return hashes.get("quickXorHash") or hashes.get("sha256Hash") or item.get("cTag") or item.get("eTag")


def _onedrive_upsert(drive_id, item, root_id):
    if "file" not in item:
        return None
    if not ONEDRIVE_RECURSIVE and (item.get("parentReference") or {}).get("id") != root_id:
        return None
    if not _allowed(item.get("name", "")):
        return None
    ref = f"{drive_id}:{item['id']}"
    upsert_source("onedrive", ref, item["name"], _onedrive_file_hash(item), item.get("webUrl"))
    return ref


def sync_onedrive():
    if not ONEDRIVE_SHARE_URLS:
        return
    if not (M365_TENANT_ID and M365_CLIENT_ID and M365_CLIENT_SECRET):
        log_error("[ingest] 已設定 ONEDRIVE_SHARE_URLS，但 M365_TENANT_ID／CLIENT_ID／CLIENT_SECRET 不完整")
        return

    for url in ONEDRIVE_SHARE_URLS:
        key = "onedrive:" + hashlib.sha1(url.encode("utf-8")).hexdigest()[:16]
        try:
            root = _graph_get(f"/shares/{_share_id(url)}/driveItem").json()
            drive_id = root["parentReference"]["driveId"]
            root_id = root["id"]

            # 分享的是單一檔案
            if "file" in root:
                _onedrive_upsert(drive_id, root, (root.get("parentReference") or {}).get("id"))
                _save_sync_state(key)
                continue

            try:
                _onedrive_delta(key, drive_id, root_id)
            except requests.HTTPError as e:
                code = e.response.status_code if e.response is not None else 0
                if code == 410:
                    # deltaLink 過期，需要重新完整同步
                    log_info(f"[ingest] OneDrive deltaLink 過期，重新完整同步：{url}")
                    _save_sync_state(key, delta_link="")
                    _onedrive_delta(key, drive_id, root_id, reset=True)
                else:
                    # 部分 OneDrive for Business／SharePoint 環境不支援子資料夾 delta → 改用完整列表比對
                    log_info(f"[ingest] OneDrive delta 不可用（HTTP {code}），改用完整列表：{url}")
                    _onedrive_full_list(key, drive_id, root_id)
        except Exception as e:
            log_error(f"[ingest] OneDrive 同步失敗（{url}）：{e}")
            _save_sync_state(key, error=str(e)[:4000])


def _onedrive_delta(key, drive_id, root_id, reset=False):
    """
    delta 查詢：第一次回傳全部項目，之後只回傳變更（含刪除）。
    deltaLink 存在 IT_KB_SyncState.DeltaLink（以 "delta:" 開頭區分完整列表模式）
    """
    saved = None if reset else _get_sync_state(key)
    url = saved[6:] if saved and saved.startswith("delta:") else f"/drives/{drive_id}/items/{root_id}/delta"
    n_items = 0
    while url:
        data = _graph_get(url).json()
        for item in data.get("value", []):
            n_items += 1
            if "deleted" in item:
                mark_deleting("onedrive", f"{drive_id}:{item['id']}")
            else:
                _onedrive_upsert(drive_id, item, root_id)
        if "@odata.nextLink" in data:
            url = data["@odata.nextLink"]
        else:
            _save_sync_state(key, delta_link="delta:" + data.get("@odata.deltaLink", ""))
            url = None
    log_info(f"[ingest] OneDrive delta 同步完成（{key}），處理 {n_items} 個項目")


def _onedrive_full_list(key, drive_id, root_id):
    """完整列表模式：列出全部檔案，與上次清單比對找出被刪除的檔案（清單存在 DeltaLink 欄位）"""
    seen, queue = set(), [root_id]
    while queue:
        fid = queue.pop()
        url = f"/drives/{drive_id}/items/{fid}/children?$top=200"
        while url:
            data = _graph_get(url).json()
            for item in data.get("value", []):
                if "folder" in item:
                    if ONEDRIVE_RECURSIVE:
                        queue.append(item["id"])
                    continue
                ref = _onedrive_upsert(drive_id, item, root_id)
                if ref:
                    seen.add(ref)
            url = data.get("@odata.nextLink")

    saved = _get_sync_state(key)
    if saved and saved.startswith("full:"):
        for ref in set(json.loads(saved[5:])) - seen:
            mark_deleting("onedrive", ref)
    _save_sync_state(key, delta_link="full:" + json.dumps(sorted(seen)))


def _fetch_onedrive(ref):
    drive_id, item_id = ref.split(":", 1)
    return _graph_get(f"/drives/{drive_id}/items/{item_id}/content").content


# =====================================================
# 9. 刪除／匯入
# =====================================================
def process_deleting():
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT DocumentID, FileName FROM IT_KB_SourceDoc WHERE Status = 'Deleting'")
        rows = cursor.fetchall()
    for r in rows:
        conn = get_mssql_conn()
        try:
            conn.autocommit = False
            cursor = conn.cursor()
            _delete_doc_rows(cursor, r.DocumentID)
            cursor.execute("DELETE FROM IT_KB_SourceDoc WHERE DocumentID = ?", r.DocumentID)
            conn.commit()
            log_info(f"[ingest] 已刪除文件 #{r.DocumentID} {r.FileName}")
        except Exception as e:
            conn.rollback()
            log_error(f"[ingest] 刪除文件 #{r.DocumentID} 失敗：{e}")
        finally:
            conn.close()
    if rows:
        _cleanup_orphan_entities()


def _fetch_bytes(doc_id, source_type, source_ref):
    if source_type == "upload":
        with get_mssql_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT FileContent FROM IT_KB_SourceDoc WHERE DocumentID = ?", doc_id)
            row = cursor.fetchone()
            if not row or row.FileContent is None:
                raise ValueError("後台上傳的文件沒有 FileContent")
            return bytes(row.FileContent)
    if source_type == "folder":
        with open(os.path.join(DOC_DROP_FOLDER, source_ref), "rb") as fp:
            return fp.read()
    if source_type == "gdrive":
        return _fetch_gdrive(source_ref)
    if source_type == "onedrive":
        return _fetch_onedrive(source_ref)
    raise ValueError(f"未知的 SourceType：{source_type}")


def process_pending(only_doc_id=None):
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        if only_doc_id:
            cursor.execute("UPDATE IT_KB_SourceDoc SET Status = 'Pending' WHERE DocumentID = ?", only_doc_id)
            conn.commit()
            cursor.execute("SELECT DocumentID, FileName, SourceType, SourceRef FROM IT_KB_SourceDoc WHERE DocumentID = ?",
                           only_doc_id)
        else:
            cursor.execute("SELECT DocumentID, FileName, SourceType, SourceRef FROM IT_KB_SourceDoc "
                           "WHERE Status = 'Pending' ORDER BY DocumentID")
        rows = cursor.fetchall()

    for r in rows:
        # 搶占：只有仍是 Pending 的才處理（避免與其他程序重複）
        with get_mssql_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE IT_KB_SourceDoc SET Status = 'Processing', UpdatedAt = SYSUTCDATETIME() "
                           "WHERE DocumentID = ? AND Status = 'Pending'", r.DocumentID)
            claimed = cursor.rowcount
            conn.commit()
        if not claimed:
            continue
        try:
            data = _fetch_bytes(r.DocumentID, r.SourceType, r.SourceRef)
            ingest_document(r.DocumentID, r.FileName, data)
        except Exception as e:
            log_error(f"[ingest] 文件 #{r.DocumentID} {r.FileName} 匯入失敗：{e}")
            with get_mssql_conn() as conn:
                cursor = conn.cursor()
                # 外部 API 的錯誤訊息可能帶金鑰，ErrorMsg 會顯示在後台，要先遮掉
                cursor.execute("UPDATE IT_KB_SourceDoc SET Status = 'Failed', ErrorMsg = ?, UpdatedAt = SYSUTCDATETIME() "
                               "WHERE DocumentID = ?", (mask_secrets(str(e))[:4000], r.DocumentID))
                conn.commit()


def reset_stuck():
    """程式中斷時留下的 Processing → 改回 Pending"""
    with get_mssql_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE IT_KB_SourceDoc SET Status = 'Pending' WHERE Status = 'Processing'")
        if cursor.rowcount:
            log_info(f"[ingest] 重設 {cursor.rowcount} 筆中斷的 Processing 文件")
        conn.commit()


def run_once():
    for name, fn in (("資料夾", scan_folder), ("Google Drive", sync_gdrive), ("OneDrive", sync_onedrive)):
        try:
            fn()
        except Exception as e:
            log_error(f"[ingest] {name} 收件失敗：{e}")
    for name, fn in (("刪除", process_deleting), ("匯入", process_pending)):
        try:
            fn()
        except Exception as e:
            log_error(f"[ingest] {name}階段失敗：{mask_secrets(str(e))}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="RBIT 文件 GraphRAG 匯入")
    parser.add_argument("--once", action="store_true", help="跑一輪就結束")
    parser.add_argument("--loop", action="store_true", help="常駐，每 INGEST_INTERVAL_SEC 秒一輪")
    parser.add_argument("--doc", type=int, help="只重新處理指定 DocumentID")
    args = parser.parse_args()

    if args.loop:
        # 常駐模式：資料庫暫時掛掉不要讓程序死掉，下一輪再試
        log_info(f"[ingest] 常駐模式，每 {INGEST_INTERVAL_SEC} 秒一輪")
        started = False
        while True:
            if check_db():
                if not started:
                    reset_stuck()
                    started = True
                run_once()
            time.sleep(INGEST_INTERVAL_SEC)

    # 單次模式：連不上就直接結束，回傳非 0 讓排程器看得出失敗
    if not check_db():
        sys.exit(1)
    reset_stuck()
    if args.doc:
        process_pending(only_doc_id=args.doc)
    else:
        run_once()
